# -*- coding: utf-8 -*-
"""
合并逻辑：文件分组、数据对齐、输出写入
"""

import time
import threading
from pathlib import Path
from typing import List, Any, Dict, Tuple, Callable, Optional
from collections import defaultdict, deque

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from .config import MODE_FAST, TABLE_DETECTION_CONFIG
from .detection import (
    detect_main_table_region, build_header_signature,
    calculate_header_similarity, count_non_empty_cells,
    looks_like_repeated_header, looks_like_page_marker,
    looks_like_print_info, looks_like_total_or_subtotal,
    score_data_row, estimate_table_column_count,
)
from .reading import (
    read_file_data, get_sheet_for_merge_info, detect_file_format,
    ExcelMergeError, FileContext,
)
from .cleaning import apply_column_cleaning


def group_files_by_structure(file_results: List[Dict]) -> Dict[str, List[Dict]]:
    """按表头结构分组文件"""
    groups: Dict[str, List[Dict]] = {}
    similarity_threshold = TABLE_DETECTION_CONFIG['signature_similarity_threshold']
    for file_result in file_results:
        if not file_result.get('success'):
            continue
        signature = file_result.get('signature', '')
        if not signature:
            continue
        matched = False
        for group_sig in groups:
            if calculate_header_similarity(signature, group_sig) >= similarity_threshold:
                groups[group_sig].append(file_result)
                matched = True
                break
        if not matched:
            groups[signature] = [file_result]
    return groups


def build_align_plan(row_columns: List[Any], unified_columns: List[Any]) -> List[Optional[int]]:
    """预计算对齐计划：返回每个统一列对应数据行的索引（None 表示该文件无此列）。

    重复列名按出现顺序一一配对；计划与具体数据行无关，可在行循环外构建一次复用。
    """
    name_to_positions = defaultdict(deque)
    for i, c in enumerate(row_columns):
        if c is not None:
            name_to_positions[str(c).strip()].append(i)
    plan: List[Optional[int]] = []
    for col in unified_columns:
        col_key = str(col).strip() if col is not None else ''
        positions = name_to_positions.get(col_key)
        plan.append(positions.popleft() if positions else None)
    return plan


def _apply_align_plan(row_data: List[Any], plan: List[Optional[int]]) -> List[Any]:
    """按计划对齐单行数据：取出对应位置的值，缺列或越界补 None。"""
    return [row_data[i] if i is not None and i < len(row_data) else None for i in plan]


def align_data_to_columns(row_data: List[Any], row_columns: List[Any], unified_columns: List[Any]) -> List[Any]:
    """将数据行按统一列顺序对齐。

    重复列名按出现顺序一一配对，避免字典映射互相覆盖导致数据错位。
    """
    return _apply_align_plan(row_data, build_align_plan(row_columns, unified_columns))


def _template_exists(header_templates: List[Dict[str, Any]], raw_header: List[Any], header_row: int) -> bool:
    """判断相同表头与表头行位置的模板是否已存在（避免重复追加相同模板）。"""
    for tpl in header_templates:
        if tpl.get('raw_header') == raw_header and tpl.get('header_row') == header_row:
            return True
    return False


def build_fallback_output_rows(file_result: Dict[str, Any]) -> Tuple[List[str], List[List[Any]]]:
    """为失败文件构建保底输出"""
    raw_data = file_result.get('raw_data') or []
    max_cols = max((len(r) for r in raw_data), default=0)
    base_headers = [f'原始列{i + 1}' for i in range(max_cols)]
    headers = base_headers + ['来源文件', '原始行号', '识别阶段失败原因', '降级路径']
    rows: List[List[Any]] = []
    fail_reason = file_result.get('error') or '未知失败'
    degradation_path = file_result.get('degradation_path') or 'fallback_raw_extract'
    file_name = file_result.get('file_name', '')
    for row_idx, row in enumerate(raw_data, start=1):
        output_row = list(row) + [None] * (max_cols - len(row))
        output_row.extend([file_name, row_idx, fail_reason, degradation_path])
        rows.append(output_row)
    if not rows:
        rows.append([None] * max_cols + [file_name, None, fail_reason, degradation_path])
    return headers, rows


def _try_match_header_template(
    data: List[List[Any]],
    templates: List[Dict[str, Any]],
) -> Optional[Dict[str, Any]]:
    """
    尝试用已有表头模板匹配当前数据

    匹配条件：data 中与模板 header_row 位置对应的行，与模板 raw_header 完全一致。
    '完全一致'指两行长度相同，且对应位置 str(cell).strip() 后完全相同。

    Args:
        data: 当前 sheet 的原始数据
        templates: 已识别成功的表头模板列表

    Returns:
        匹配成功的模板（字典），未匹配返回 None
    """
    if not data or not templates:
        return None

    # 从配置读取检查行数上限，避免不必要的遍历
    check_limit = min(len(data), TABLE_DETECTION_CONFIG.get('template_check_limit', 5))

    for tpl in templates:
        hdr_row_idx = tpl.get('header_row', -1)
        raw_header = tpl.get('raw_header')
        if hdr_row_idx < 0 or not raw_header:
            continue
        # header_row 是 0-based 索引，必须在可检查范围内
        if hdr_row_idx >= check_limit:
            continue

        actual_row = data[hdr_row_idx]
        # 长度必须一致
        if len(actual_row) != len(raw_header):
            continue

        # 逐格比较
        matched = True
        for i in range(len(raw_header)):
            val_actual = str(actual_row[i]).strip() if i < len(actual_row) else ''
            val_tpl = str(raw_header[i]).strip() if i < len(raw_header) else ''
            if val_actual != val_tpl:
                matched = False
                break

        if matched:
            return tpl

    return None


def process_single_sheet(
    file_path: str,
    file_name: str,
    file_ext: str,
    sheet_name: Optional[str],
    is_fast_mode: bool,
    log_callback: Callable[[str], None],
    file_ctx: Optional[FileContext] = None,
    header_templates: Optional[List[Dict[str, Any]]] = None,
    cancel_event: Optional[threading.Event] = None,
) -> Dict[str, Any]:
    """
    处理单个文件的单个sheet

    Args:
        file_path: 文件路径
        file_name: 文件名
        file_ext: 文件扩展名
        sheet_name: sheet名称（None表示第一个sheet）
        is_fast_mode: 是否快速模式
        log_callback: 日志回调
        file_ctx: 文件级上下文（提供格式缓存和工作簿复用，为None时回退到原有行为）
        header_templates: 已识别成功的表头模板列表，用于跳过重复的表头识别

    Returns:
        文件处理结果字典
    """
    sheet_display = sheet_name or "(默认)"
    file_result: Dict[str, Any] = {
        'file_path': file_path,
        'file_name': file_name,
        'sheet_name': sheet_name,
        'display_name': f"{file_name} [{sheet_display}]" if sheet_name else file_name,
        'file_ext': file_ext,
        'success': False,
        'is_encrypted': False,
        'header_row': -1,
        'header': [],
        'compact_header': [],
        'data_rows': [],
        'compact_data_rows': [],
        'col_count': 0,
        'signature': '',
        'error': None,
        'source_wb': None,
        'source_ws': None,
        'raw_data': [],
        'audit': {},
        'degradation_path': '',
        'timing': {},
    }

    try:
        data, max_row, max_col, source_wb, source_ws, is_encrypted, error_msg, read_timing = read_file_data(
            file_path, sheet_name=sheet_name, is_fast_mode=is_fast_mode
        )

        if is_encrypted:
            file_result['is_encrypted'] = True
            file_result['error'] = "文件受密码保护"
            log_callback(f"  ⚠️ [{sheet_display}] 受密码保护，已跳过")
        elif data is None:
            file_result['error'] = error_msg or "读取失败"
            file_result['degradation_path'] = 'read_failed'
            log_callback(f"  ❌ [{sheet_display}] 读取失败: {error_msg}")
        else:
            file_result['raw_data'] = data
            file_result['timing']['read'] = read_timing

            # 获取合并单元格信息（仅xlsx）
            # 优先使用 FileContext（复用工作簿、缓存格式），否则回退到原有行为
            merge_info_wb = None
            merge_info_sheet = None
            if file_ctx is not None:
                actual_format = file_ctx.file_format
            else:
                actual_format = detect_file_format(file_path)
            if actual_format == 'xlsx':
                if file_ctx is not None:
                    merge_info_wb, merge_info_sheet = file_ctx.get_sheet_for_merge_info(sheet_name)
                else:
                    merge_info_wb, merge_info_sheet = get_sheet_for_merge_info(file_path, sheet_name)

            # 尝试匹配已有表头模板，跳过重复的表头识别
            matched_tpl = _try_match_header_template(data, header_templates or [])

            try:
                if matched_tpl is not None:
                    # 模板命中：直接复用表头信息，仍需执行数据行扩展
                    log_callback(f"  [{sheet_display}] 表头模板命中，跳过表头识别")

                    header_row = matched_tpl['header_row']
                    header_row_count = matched_tpl.get('header_row_count', 1)
                    header = matched_tpl['raw_header']
                    compact_header = matched_tpl['compact_header']
                    valid_columns = matched_tpl.get('valid_columns', [])

                    # 数据行扩展：从 header_row 之后开始，带终止检测
                    # 逻辑与 detect_main_table_region 中 _try_detect_in_block 的数据行扩展一致
                    data_start_offset = header_row + header_row_count
                    # 限制扫描上界，避免对整个文件做无意义遍历
                    max_scan_rows = max(matched_tpl.get('data_row_count', 100) * 3, 100)
                    scan_end = min(len(data), data_start_offset + max_scan_rows)
                    max_cf = TABLE_DETECTION_CONFIG['max_consecutive_failures']
                    min_table_rows = TABLE_DETECTION_CONFIG['min_table_rows']
                    data_threshold = TABLE_DETECTION_CONFIG['data_row_score_threshold']

                    # 估算期望列数
                    sample_rows = [data[r] for r in range(data_start_offset, min(data_start_offset + 15, scan_end))]
                    expected_col_count = estimate_table_column_count(header, sample_rows)

                    # 构建合并表头：优先使用模板中存储的多行合并表头
                    merged_header = matched_tpl.get('merged_header', header)

                    data_rows = []
                    consecutive_failures = 0
                    data_start_row = -1

                    for row_idx in range(data_start_offset, scan_end):
                        # 每 100 行检查一次取消信号，避免大文件处理时取消响应延迟
                        if cancel_event is not None and row_idx % 100 == 0 and cancel_event.is_set():
                            log_callback("  合并已被用户取消，终止当前文件处理")
                            break

                        row = data[row_idx]

                        # 快速预过滤：空行直接跳过，不终止采集。
                        # 与 detect_main_table_region 保持一致：空行只是分隔符，
                        # 仅当连续多个“非空但非数据”的行时才终止。
                        non_empty_count = count_non_empty_cells(row)
                        if non_empty_count == 0:
                            continue

                        # 快速预过滤：噪声行
                        if looks_like_repeated_header(row, merged_header):
                            consecutive_failures += 0.5
                            continue
                        if looks_like_page_marker(row) or looks_like_print_info(row):
                            consecutive_failures += 0.5
                            continue
                        if looks_like_total_or_subtotal(row):
                            consecutive_failures += 0.5
                            continue

                        # 通过预过滤的行做完整评分
                        sc, row_type = score_data_row(row, merged_header, expected_col_count)
                        if sc >= data_threshold:
                            data_rows.append(row)
                            if data_start_row == -1:
                                data_start_row = row_idx
                            consecutive_failures = 0
                        else:
                            consecutive_failures += 1
                            if consecutive_failures >= max_cf and len(data_rows) >= min_table_rows:
                                break

                    # 构建 compact_data_rows：只保留 valid_columns 对应的列
                    compact_data_rows = []
                    if valid_columns:
                        for row in data_rows:
                            compact_row = [row[c] if c < len(row) else None for c in valid_columns]
                            compact_data_rows.append(compact_row)
                    else:
                        compact_data_rows = [list(r) for r in data_rows]

                    file_result['success'] = True
                    file_result['header_row'] = header_row
                    file_result['header_row_count'] = header_row_count
                    file_result['header'] = header
                    file_result['compact_header'] = compact_header
                    file_result['valid_columns'] = valid_columns
                    file_result['data_rows'] = data_rows
                    file_result['compact_data_rows'] = compact_data_rows
                    file_result['col_count'] = len(compact_header)
                    file_result['signature'] = build_header_signature(compact_header)
                    file_result['degradation_path'] = 'header_template_reuse'
                    # 补充审计信息和数据起止行，与原路径保持一致
                    file_result['audit'] = {
                        'degradation_path': 'header_template_reuse',
                        'matched_template_header_row': header_row,
                        'data_start': data_start_row,
                        'data_end': data_start_row + len(data_rows) - 1 if data_start_row >= 0 else -1,
                    }
                    file_result['data_start'] = data_start_row
                    file_result['data_end'] = data_start_row + len(data_rows) - 1 if data_start_row >= 0 else -1

                    log_callback(f"  [{sheet_display}] 表头行: {header_row + 1}, 数据行: {len(data_rows)}, 列数: {len(compact_header)}")
                else:
                    # 未命中模板：走完整的表头识别流程
                    table_result = detect_main_table_region(data, sheet=merge_info_sheet)

                    if table_result['success']:
                        file_result['success'] = True
                        file_result['header_row'] = table_result['header_row']
                        file_result['header'] = table_result['header']
                        file_result['compact_header'] = table_result['compact_header']
                        file_result['merged_header'] = table_result.get('merged_header', table_result['header'])
                        file_result['data_rows'] = table_result['data_rows']
                        file_result['compact_data_rows'] = table_result['compact_data_rows']
                        file_result['col_count'] = table_result['col_count']
                        file_result['signature'] = build_header_signature(table_result['compact_header'])
                        file_result['audit'] = table_result.get('audit', {})
                        file_result['degradation_path'] = table_result.get('audit', {}).get('degradation_path', 'primary_block')
                        file_result['timing']['detection'] = table_result.get('timing', {})
                        log_callback(f"  [{sheet_display}] 表头行: {table_result['header_row'] + 1}, 数据行: {len(table_result['data_rows'])}, 列数: {len(table_result['compact_header'])}")
                    else:
                        file_result['error'] = table_result['error']
                        file_result['audit'] = table_result.get('audit', {})
                        file_result['degradation_path'] = table_result.get('audit', {}).get('degradation_path', 'fallback_raw_extract')
                        log_callback(f"  [{sheet_display}] 失败 - {table_result['error']}")
            finally:
                # 仅在未使用 FileContext 时关闭工作簿（FileContext 的工作簿由外部统一管理）
                if file_ctx is None and merge_info_wb:
                    try:
                        merge_info_wb.close()
                    except Exception:
                        pass

    except Exception as e:
        file_result['error'] = str(e)
        log_callback(f"  [{sheet_display}] 异常: {e}")

    return file_result


def merge_excel_files_with_callback(
    file_list: List[str],
    output_path: str,
    mode: str,
    log_callback: Callable[[str], None],
    progress_callback: Callable[[int, int], None],
    file_sheet_mapping: Optional[Dict[str, List[str]]] = None,
    cancel_event: Optional[threading.Event] = None,
) -> Dict[str, Any]:
    """
    合并多个Excel文件（带回调的GUI版本）

    Args:
        file_list: 要合并的Excel文件路径列表
        output_path: 输出文件路径
        mode: 处理模式
        log_callback: 日志回调
        progress_callback: 进度回调(current, total)
        file_sheet_mapping: 文件路径 -> 要处理的sheet名称列表，为None时只处理第一个sheet
        cancel_event: 取消事件，外部设置后中断合并流程
    """
    # 安全日志回调：防止 GUI 关闭等场景下异常中断合并流程
    def _safe_log(msg):
        try:
            log_callback(msg)
        except Exception:
            pass

    # 安全进度回调：防止 GUI 关闭时 progress_callback 异常中断合并
    def _safe_progress(current, total):
        try:
            progress_callback(current, total)
        except Exception:
            pass

    if not file_list:
        raise ExcelMergeError("没有可合并的文件")

    total_files = len(file_list)
    is_fast_mode = (mode == MODE_FAST)

    # 构建处理任务列表：每个文件的每个sheet都是一个任务
    tasks: List[Tuple[str, str, str, Optional[str]]] = []  # (file_path, file_name, file_ext, sheet_name)
    for file_path in file_list:
        file_name = Path(file_path).name
        file_ext = Path(file_path).suffix.lower()
        sheets = file_sheet_mapping.get(file_path) if file_sheet_mapping else None
        if sheets:
            for s in sheets:
                tasks.append((file_path, file_name, file_ext, s))
        else:
            tasks.append((file_path, file_name, file_ext, None))

    total_tasks = len(tasks)
    file_results: List[Dict[str, Any]] = []
    format_stats: Dict[str, int] = {'xlsx': 0, 'xls': 0, 'csv': 0}
    encrypted_files: List[str] = []

    _safe_log("=" * 50)
    _safe_log("阶段1: 读取文件并检测表格区域")
    _safe_log("=" * 50)
    _safe_log("")

    # 为每个唯一文件创建 FileContext，缓存格式和工作簿
    file_contexts: Dict[str, FileContext] = {}

    # 表头模板列表：首个成功识别的表头会作为模板，后续同结构文件可复用
    header_templates: List[Dict[str, Any]] = []

    try:
        for idx, (file_path, file_name, file_ext, sheet_name) in enumerate(tasks, 1):
            # 检查取消事件
            if cancel_event is not None and cancel_event.is_set():
                _safe_log("合并已被用户取消")
                break

            _safe_progress(idx, total_tasks)
            display = f"{file_name} [{sheet_name}]" if sheet_name else file_name
            _safe_log(f"[{idx}/{total_tasks}] 读取: {display}")

            # 获取或创建该文件的 FileContext
            if file_path not in file_contexts:
                file_contexts[file_path] = FileContext(file_path)
            ctx = file_contexts[file_path]

            t0_task = time.perf_counter()
            result = process_single_sheet(
                file_path, file_name, file_ext, sheet_name,
                is_fast_mode, _safe_log,
                file_ctx=ctx,
                header_templates=header_templates,
                cancel_event=cancel_event,
            )
            elapsed_task = time.perf_counter() - t0_task

            # 输出分段耗时
            timings = ctx.get_timings()
            if timings:
                timing_parts = [f"{k}={v:.3f}s" for k, v in timings.items()]
                _safe_log(f"  ⏱ 文件缓存计时: {', '.join(timing_parts)}")
            _safe_log(f"  ⏱ 本任务总耗时: {elapsed_task:.3f}s")

            if result.get('is_encrypted'):
                encrypted_files.append(display)
            elif result.get('success'):
                if file_ext == '.xlsx':
                    format_stats['xlsx'] += 1
                elif file_ext == '.xls':
                    format_stats['xls'] += 1
                else:
                    format_stats['csv'] += 1

                # 将成功识别的表头信息追加到模板列表，供后续同结构文件复用
                # 相同表头（raw_header 与 header_row 均相同）只保留一个模板，避免重复追加
                if not _template_exists(header_templates, result['header'], result['header_row']):
                    header_templates.append({
                        'raw_header': result['header'],
                        'header_row': result['header_row'],
                        'header_row_count': result.get('header_row_count', 1),
                        'compact_header': result['compact_header'],
                        'valid_columns': result.get('valid_columns', []),
                        'data_row_count': len(result.get('compact_data_rows', [])),
                        'merged_header': result.get('merged_header', result['header']),
                    })

            file_results.append(result)
            _safe_log("")
    finally:
        # 统一关闭所有 FileContext（释放工作簿资源），即使循环中异常也确保执行
        for fp, ctx in file_contexts.items():
            ctx.close()

    # 如果用户已取消，直接返回（不执行阶段2/3）
    if cancel_event is not None and cancel_event.is_set():
        _safe_log("合并已被用户取消，跳过后续阶段")
        return {
            'total_files': total_tasks,
            'total_rows': 0,
            'format_stats': format_stats,
            'groups': {},
            'success_count': 0,
            'encrypted_count': len(encrypted_files),
            'failed_count': 0,
            'encrypted_files': encrypted_files,
            'failed_files': [],
            'output_path': output_path,
            'fallback_rows': 0,
            'audit': {},
            'cancelled': True,
        }

    # 阶段2: 分组
    _safe_log("=" * 50)
    _safe_log("阶段2: 按结构分组")
    _safe_log("=" * 50)
    _safe_log("")

    groups = group_files_by_structure(file_results)
    _safe_log(f"发现 {len(groups)} 个结构组:")
    for idx, (sig, files) in enumerate(groups.items(), 1):
        sig_display = sig[:50] + '...' if len(sig) > 50 else sig
        _safe_log(f"  组{idx} ({len(files)}个): {sig_display}")
    _safe_log("")

    success_count = sum(1 for f in file_results if f.get('success'))
    encrypted_count = len(encrypted_files)
    failed_count = sum(1 for f in file_results if not f.get('success') and not f.get('is_encrypted'))

    _safe_log(f"任务总数: {total_tasks}")
    _safe_log(f"  成功: {success_count}")
    _safe_log(f"  加密: {encrypted_count}")
    _safe_log(f"  失败: {failed_count}")
    _safe_log("")

    # 阶段3: 合并
    _safe_log("=" * 50)
    _safe_log("阶段3: 合并数据")
    _safe_log("=" * 50)
    _safe_log("")

    all_files_to_merge = []
    for group_sig, group_files in groups.items():
        _safe_log(f"组 [{group_sig[:40]}...]: {len(group_files)}个")
        for f in group_files:
            _safe_log(f"  - {f['display_name']} ({len(f.get('compact_data_rows', []))}行)")
        all_files_to_merge.extend(group_files)
    _safe_log(f"共需合并: {len(all_files_to_merge)}个")
    _safe_log("")

    # 收集统一列（重复列名按各文件最大出现次数保留，同名列数据不丢）
    max_counts: Dict[str, int] = {}
    for f in all_files_to_merge:
        local_counts: Dict[str, int] = {}
        for col in f.get('compact_header', []):
            if col is None:
                continue
            col_str = str(col).strip()
            if col_str:
                local_counts[col_str] = local_counts.get(col_str, 0) + 1
        for name, cnt in local_counts.items():
            if cnt > max_counts.get(name, 0):
                max_counts[name] = cnt

    unified_columns = []
    added_counts: Dict[str, int] = {}
    for f in all_files_to_merge:
        for col in f.get('compact_header', []):
            if col is not None:
                col_str = str(col).strip()
                if col_str and added_counts.get(col_str, 0) < max_counts[col_str]:
                    unified_columns.append(col_str)
                    added_counts[col_str] = added_counts.get(col_str, 0) + 1

    _safe_log(f"统一列数: {len(unified_columns)}")
    _safe_log(f"列名: {unified_columns}")
    _safe_log("")

    fallback_files = [f for f in file_results if (not f.get('success')) and (not f.get('is_encrypted'))]
    fallback_row_count = 0
    total_data_rows = 0
    has_main_data = len(all_files_to_merge) > 0 and len(unified_columns) > 0

    if is_fast_mode:
        _safe_log("使用快速模式（Write-Only）写入...")
        result_wb = Workbook(write_only=True)
        if has_main_data:
            result_ws = result_wb.create_sheet("合并数据")
            result_ws.append(list(unified_columns) + ['数据来源'])
            for file_result in all_files_to_merge:
                display_name = file_result['display_name']
                compact_header = file_result['compact_header']
                _safe_log(f"合并: {display_name} ({len(file_result['compact_data_rows'])}行)")
                clean_rows = apply_column_cleaning(compact_header, file_result['compact_data_rows'])
                align_plan = build_align_plan(compact_header, unified_columns)
                for row_data in clean_rows:
                    aligned_row = _apply_align_plan(row_data, align_plan)
                    result_ws.append(aligned_row + [display_name])
                    total_data_rows += 1
        if fallback_files:
            fallback_ws = result_wb.create_sheet("保底数据")
            fallback_headers_written = False
            for file_result in fallback_files:
                headers, rows_data = build_fallback_output_rows(file_result)
                if not fallback_headers_written:
                    fallback_ws.append(headers)
                    fallback_headers_written = True
                for row_data in rows_data:
                    fallback_ws.append(row_data)
                fallback_row_count += len(rows_data)
        result_wb.save(output_path)
    else:
        _safe_log("使用完整模式（保留样式）写入...")
        result_wb = Workbook()
        col_max_width: Dict[int, int] = {}
        if has_main_data:
            result_ws = result_wb.active
            result_ws.title = "合并数据"
            current_row = 1
            for col_idx, col_name in enumerate(unified_columns, 1):
                cell = result_ws.cell(row=current_row, column=col_idx)
                cell.value = col_name
                cell.font = Font(bold=True)
                col_max_width[col_idx] = max(col_max_width.get(col_idx, 0), len(str(col_name)))
            source_col = result_ws.cell(row=current_row, column=len(unified_columns) + 1)
            source_col.value = '数据来源'
            source_col.font = Font(bold=True)
            source_col.alignment = Alignment(horizontal='center')
            current_row += 1

            for file_result in all_files_to_merge:
                display_name = file_result['display_name']
                compact_header = file_result['compact_header']
                _safe_log(f"合并: {display_name} ({len(file_result['compact_data_rows'])}行)")
                clean_rows = apply_column_cleaning(compact_header, file_result['compact_data_rows'])
                align_plan = build_align_plan(compact_header, unified_columns)
                for row_data in clean_rows:
                    aligned_row = _apply_align_plan(row_data, align_plan)
                    for col_idx, value in enumerate(aligned_row, 1):
                        cell = result_ws.cell(row=current_row, column=col_idx)
                        cell.value = value
                        if value is not None:
                            col_max_width[col_idx] = max(col_max_width.get(col_idx, 0), len(str(value)))
                    source_cell = result_ws.cell(row=current_row, column=len(unified_columns) + 1)
                    source_cell.value = display_name
                    source_cell.alignment = Alignment(horizontal='left')
                    current_row += 1
                    total_data_rows += 1

            for col_idx, max_width in col_max_width.items():
                result_ws.column_dimensions[get_column_letter(col_idx)].width = min(max_width + 2, 50)
            result_ws.column_dimensions[get_column_letter(len(unified_columns) + 1)].width = 25

        if fallback_files:
            fallback_ws = result_wb.create_sheet("保底数据") if has_main_data else result_wb.active
            if not has_main_data:
                fallback_ws.title = "保底数据"
            fallback_headers_written = False
            fallback_row_cursor = 2
            for file_result in fallback_files:
                headers, rows_data = build_fallback_output_rows(file_result)
                if not fallback_headers_written:
                    for col_idx, col_name in enumerate(headers, 1):
                        cell = fallback_ws.cell(row=1, column=col_idx)
                        cell.value = col_name
                        cell.font = Font(bold=True)
                    fallback_headers_written = True
                    fallback_row_cursor = 2
                for row_data in rows_data:
                    for col_idx, value in enumerate(row_data, 1):
                        fallback_ws.cell(row=fallback_row_cursor, column=col_idx).value = value
                    fallback_row_cursor += 1
                fallback_row_count += len(rows_data)
        result_wb.save(output_path)

    _safe_log("")
    _safe_log("=" * 50)
    _safe_log("合并完成!")
    _safe_log("=" * 50)
    _safe_log(f"输出文件: {Path(output_path).name}")
    _safe_log(f"总数据行数: {total_data_rows}")
    if fallback_row_count > 0:
        _safe_log(f"保底输出行数: {fallback_row_count}")

    return {
        'total_files': total_tasks,
        'total_rows': total_data_rows,
        'format_stats': format_stats,
        'groups': {sig: [f['display_name'] for f in files] for sig, files in groups.items()},
        'success_count': success_count,
        'encrypted_count': encrypted_count,
        'failed_count': failed_count,
        'encrypted_files': encrypted_files,
        'failed_files': [f['display_name'] for f in file_results if not f.get('success') and not f.get('is_encrypted')],
        'output_path': output_path,
        'fallback_rows': fallback_row_count,
        'audit': {f['display_name']: f.get('audit', {}) for f in file_results},
        'cancelled': False,
    }
