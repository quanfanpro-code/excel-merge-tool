# -*- coding: utf-8 -*-
"""端到端集成测试：验证金额/日期清洗真正写入合并结果"""
import csv
import os
import shutil
import tempfile

from openpyxl import load_workbook

from core.merging import merge_excel_files_with_callback


def _write_csv(path, rows):
    with open(path, 'w', encoding='utf-8', newline='') as f:
        writer = csv.writer(f)
        writer.writerows(rows)


def _noop(*args, **kwargs):
    pass


def test_merge_output_contains_cleaned_amount_and_date():
    """合并两个同表头 CSV，输出文件的金额列应为数字、日期列应统一为 YYYY-MM-DD"""
    tmpdir = tempfile.mkdtemp()
    try:
        _write_csv(os.path.join(tmpdir, 'a.csv'), [
            ['交易日期', '摘要', '金额'],
            ['20240101', '吃饭', '1,234.56'],
            ['2024/1/2', '打车', '￥50.00'],
        ])
        _write_csv(os.path.join(tmpdir, 'b.csv'), [
            ['交易日期', '摘要', '金额'],
            ['2024年1月3日', '买东西', '(100.00)'],
        ])
        output = os.path.join(tmpdir, 'result.xlsx')
        result = merge_excel_files_with_callback(
            file_list=[os.path.join(tmpdir, 'a.csv'), os.path.join(tmpdir, 'b.csv')],
            output_path=output,
            mode='fast',
            log_callback=_noop,
            progress_callback=_noop,
        )
        assert result['success_count'] == 2

        wb = load_workbook(output, read_only=True)
        ws = wb['合并数据']
        rows = list(ws.iter_rows(values_only=True))
        wb.close()

        # 第一行是表头
        assert rows[0][0] == '交易日期'
        assert rows[0][2] == '金额'

        # 金额列被清洗成可计算数字
        amounts = [r[2] for r in rows[1:]]
        assert 1234.56 in amounts          # 1,234.56
        assert 50 in amounts               # ￥50.00 -> 50
        assert -100 in amounts             # (100.00) -> -100

        # 日期列统一为 YYYY-MM-DD
        dates = [r[0] for r in rows[1:]]
        assert '2024-01-01' in dates
        assert '2024-01-02' in dates
        assert '2024-01-03' in dates
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


def test_merge_full_mode_also_cleans():
    """完整模式同样应清洗金额/日期"""
    tmpdir = tempfile.mkdtemp()
    try:
        _write_csv(os.path.join(tmpdir, 'c.csv'), [
            ['日期', '金额'],
            ['2024/3/5', '2,000.00'],
        ])
        output = os.path.join(tmpdir, 'result_full.xlsx')
        result = merge_excel_files_with_callback(
            file_list=[os.path.join(tmpdir, 'c.csv')],
            output_path=output,
            mode='full',
            log_callback=_noop,
            progress_callback=_noop,
        )
        assert result['success_count'] == 1

        wb = load_workbook(output, read_only=True)
        ws = wb['合并数据']
        rows = list(ws.iter_rows(values_only=True))
        wb.close()

        assert rows[1][0] == '2024-03-05'
        assert rows[1][1] == 2000.0
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


def test_cancel_event_returns_cancelled_flag():
    """预设取消信号时，返回结果必须带 cancelled=True（R2-4）"""
    import threading
    tmpdir = tempfile.mkdtemp()
    try:
        _write_csv(os.path.join(tmpdir, 'x.csv'), [
            ['日期', '金额'],
            ['2024/1/1', '100'],
        ])
        ev = threading.Event()
        ev.set()
        result = merge_excel_files_with_callback(
            file_list=[os.path.join(tmpdir, 'x.csv')],
            output_path=os.path.join(tmpdir, 'never.xlsx'), mode='fast',
            log_callback=_noop, progress_callback=_noop,
            cancel_event=ev,
        )
        assert result.get('cancelled') is True
        assert not os.path.exists(os.path.join(tmpdir, 'never.xlsx'))
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)
