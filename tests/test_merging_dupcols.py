# -*- coding: utf-8 -*-
"""重复列名场景：统一列必须保留重复列名，同名列数据不丢不串（R2-1/R2-1b/R2-5）"""
import csv
import os
import shutil
import tempfile

from openpyxl import load_workbook

from core.merging import (
    merge_excel_files_with_callback, align_data_to_columns, build_align_plan,
)


def _write_csv(path, rows):
    with open(path, 'w', encoding='utf-8', newline='') as f:
        csv.writer(f).writerows(rows)


def _noop(*args, **kwargs):
    pass


def test_duplicate_header_columns_all_preserved():
    """表头有两个「金额」列的文件，合并后两列数据都完整保留（R2-1）"""
    tmpdir = tempfile.mkdtemp()
    try:
        _write_csv(os.path.join(tmpdir, 'dup.csv'), [
            ['日期', '金额', '金额'],
            ['2024/1/1', '100', '200'],
            ['2024/1/2', '300', '400'],
        ])
        output = os.path.join(tmpdir, 'out.xlsx')
        result = merge_excel_files_with_callback(
            file_list=[os.path.join(tmpdir, 'dup.csv')],
            output_path=output, mode='fast',
            log_callback=_noop, progress_callback=_noop,
        )
        assert result['success_count'] == 1

        wb = load_workbook(output, read_only=True)
        rows = list(wb['合并数据'].iter_rows(values_only=True))
        wb.close()

        # 统一列必须含两个「金额」列，两条数据的两个金额值都在且不串列
        header = rows[0]
        assert list(header[:3]) == ['日期', '金额', '金额']
        data_rows = [list(r[:3]) for r in rows[1:]]
        assert ['2024-01-01', 100, 200] in data_rows
        assert ['2024-01-02', 300, 400] in data_rows
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


def test_mixed_dup_and_single_amount_aligns_in_order():
    """A 文件 2 个金额列、B 文件 1 个金额列同组：B 占第一个金额位，第二位为空"""
    tmpdir = tempfile.mkdtemp()
    try:
        _write_csv(os.path.join(tmpdir, 'a.csv'), [
            ['日期', '金额', '金额'],
            ['2024/1/1', '100', '200'],
        ])
        _write_csv(os.path.join(tmpdir, 'b.csv'), [
            ['日期', '金额'],
            ['2024/1/2', '999'],
        ])
        output = os.path.join(tmpdir, 'out.xlsx')
        result = merge_excel_files_with_callback(
            file_list=[os.path.join(tmpdir, 'a.csv'), os.path.join(tmpdir, 'b.csv')],
            output_path=output, mode='fast',
            log_callback=_noop, progress_callback=_noop,
        )
        assert result['success_count'] == 2

        wb = load_workbook(output, read_only=True)
        rows = list(wb['合并数据'].iter_rows(values_only=True))
        wb.close()

        header = list(rows[0][:3])
        assert header == ['日期', '金额', '金额']
        b_row = [r for r in rows[1:] if r[0] == '2024-01-02'][0]
        assert b_row[1] == 999 and b_row[2] is None
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


def test_build_align_plan_matches_align_data_to_columns():
    """build_align_plan 预计算结果必须与 align_data_to_columns 逐行结果一致（R2-1b 重构安全）"""
    row_columns = ['金额', '金额', '日期', '备注']
    unified = ['金额', '金额', '日期', '不存在列']
    plan = build_align_plan(row_columns, unified)
    assert plan == [0, 1, 2, None]
    row_data = [10, 20, '2024-01-01', 'x']
    via_plan = [row_data[i] if i is not None and i < len(row_data) else None for i in plan]
    assert via_plan == align_data_to_columns(row_data, row_columns, unified)
    # 短行越界安全
    assert align_data_to_columns([10], row_columns, unified) == [10, None, None, None]


def test_template_exists_dedup_check():
    """相同 raw_header + header_row 的模板必须被识别为已存在（R2-5）"""
    from core.merging import _template_exists
    templates = [{'raw_header': ['日期', '金额'], 'header_row': 0}]
    assert _template_exists(templates, ['日期', '金额'], 0) is True
    assert _template_exists(templates, ['日期', '金额'], 1) is False
    assert _template_exists(templates, ['日期', '摘要'], 0) is False


def test_same_header_files_still_hit_template():
    """3 个同表头文件：第 2、3 个仍命中模板（去重后模板可用）"""
    tmpdir = tempfile.mkdtemp()
    try:
        logs = []
        for name in ('s1.csv', 's2.csv', 's3.csv'):
            _write_csv(os.path.join(tmpdir, name), [
                ['日期', '金额'],
                ['2024/1/1', '100'],
            ])
        merge_excel_files_with_callback(
            file_list=[os.path.join(tmpdir, n) for n in ('s1.csv', 's2.csv', 's3.csv')],
            output_path=os.path.join(tmpdir, 'out.xlsx'), mode='fast',
            log_callback=logs.append, progress_callback=_noop,
        )
        hits = [m for m in logs if '表头模板命中' in m]
        assert len(hits) == 2
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)
